"""Read a seller T-12, tag every GL line with a model code, and check the result.

Usage:
    python -m valuation.t12 <t12.xls|xlsx> [--sheet T-12] [--map mappings/onesite_gl_map.csv]
                            [--out out_dir] [--no-map]

Outputs (in --out):
    t12_tagged.csv     one row per GL line: account, name, section, 12-mo total, code, how tagged
    t12_for_model.xlsx a 'T-12' sheet in the layout the model expects
                       (col A account name, cols B-M months, col N 12-mo total, col O code)
    summary.json       model totals per code, NOI check, untagged lines

Exit code 0 only if every non-zero line is tagged AND the model NOI matches the T-12 NOI.
"""
from __future__ import annotations

import argparse
import csv
import json
import os
import re
import sys
from dataclasses import dataclass, field, asdict

from .codes import MODEL_CODES, INCOME_CODES
from .rules import tag_by_keyword

GL_RE = re.compile(r"^\s*(\d{3,6}(?:[-.]\d{1,6})?)\s*[-–:]?\s*(.+?)\s*$")
TOLERANCE = 1.00  # dollars


@dataclass
class Line:
    row: int
    account: str
    name: str
    section: str  # "income" | "expense"
    months: list = field(default_factory=list)
    total: float = 0.0
    code: str | None = None
    tagged_by: str = ""


# --------------------------------------------------------------------------- reading
def _read_grid(path: str, sheet: str | None):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".xls":
        import xlrd
        book = xlrd.open_workbook(path)
        sh = _pick_sheet(book.sheet_names(), sheet, lambda n: book.sheet_by_name(n))
        return [[sh.cell_value(r, c) for c in range(sh.ncols)] for r in range(sh.nrows)]
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = _pick_sheet(wb.sheetnames, sheet, lambda n: wb[n])
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _pick_sheet(names, wanted, get):
    if wanted:
        return get(wanted)
    for n in names:
        if re.search(r"t-?12|twelve|trailing|income statement", n, re.I):
            return get(n)
    return get(names[0])


def _num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace(",", "").replace("$", "").strip()
    if re.fullmatch(r"\(.*\)", s):
        s = "-" + s[1:-1]
    try:
        return float(s)
    except ValueError:
        return None


def parse_t12(grid):
    """Return (lines, t12_noi, total_col)."""
    # 12-month total column: header says Summary/Total/YTD; else right-most numeric column
    total_col = None
    for r in grid[:20]:
        for c, v in enumerate(r):
            if isinstance(v, str) and re.fullmatch(r"\s*(summary|total|ytd|t-?12 total|12 month total)\s*", v, re.I):
                total_col = c
    if total_col is None:
        total_col = max((c for r in grid for c, v in enumerate(r) if _num(v) is not None), default=None)

    lines, noi, section = [], None, "income"
    for i, r in enumerate(grid):
        label = next((str(v).strip() for v in r[:3] if isinstance(v, str) and v.strip()), "")
        if not label:
            continue
        up = label.upper()
        if re.search(r"NET OPERATING INCOME|^NOI\b", up):
            noi = _num(r[total_col]) if total_col is not None and total_col < len(r) else None
            continue
        if re.fullmatch(r"\s*(OPERATING\s+)?EXPENSES?\s*", up) or up.startswith("TOTAL INCOME"):
            section = "expense"
            continue
        if up.startswith("TOTAL") or up.startswith("NET "):
            continue
        tot = _num(r[total_col]) if total_col is not None and total_col < len(r) else None
        if tot is None:
            continue  # section header
        m = GL_RE.match(label)
        account, name = (m.group(1), m.group(2)) if m else ("", label)
        months = [_num(v) or 0.0 for v in r[1:total_col]] if total_col else []
        lines.append(Line(i, account, name.lstrip("- ").strip(), section, months, tot))
    return lines, noi, total_col


# --------------------------------------------------------------------------- tagging
def load_map(path):
    m = {}
    if path and os.path.exists(path):
        with open(path, newline="") as f:
            for row in csv.DictReader(f):
                m[row["account"].strip()] = row["code"].strip()
    return m


def tag_lines(lines, gl_map):
    for ln in lines:
        if ln.account and ln.account in gl_map:
            ln.code, ln.tagged_by = gl_map[ln.account], "gl_map"
            continue
        code, pat = tag_by_keyword(ln.name, ln.section)
        if code:
            ln.code, ln.tagged_by = code, f"keyword:{pat.split('|')[0]}"
    return lines


def model_amount(ln: Line) -> float:
    """Amount as the model wants it: income codes positive; every other code a positive cost."""
    if ln.code in INCOME_CODES:
        return ln.total if ln.section == "income" else -ln.total
    return -ln.total if ln.section == "income" else ln.total


def summarize(lines, t12_noi):
    totals = {c: 0.0 for c in MODEL_CODES}
    for ln in lines:
        if ln.code:
            totals[ln.code] += model_amount(ln)
    income = sum(totals[c] for c in INCOME_CODES)
    costs = sum(v for c, v in totals.items() if c not in INCOME_CODES)
    model_noi = income - costs
    untagged = [ln for ln in lines if not ln.code and abs(ln.total) > 0.005]
    noi_ok = t12_noi is not None and abs(model_noi - t12_noi) <= TOLERANCE
    return {
        "lines": len(lines),
        "tagged": sum(1 for ln in lines if ln.code),
        "untagged_nonzero": [{"account": l.account, "name": l.name, "total": round(l.total, 2)} for l in untagged],
        "t12_noi": t12_noi,
        "model_noi_before_reserves": round(model_noi, 2),
        "noi_difference": None if t12_noi is None else round(model_noi - t12_noi, 2),
        "noi_check": "PASS" if noi_ok else "FAIL",
        "scenario0_t12_column": {c: round(v, 2) for c, v in totals.items() if abs(v) > 0.005},
    }


# --------------------------------------------------------------------------- output
def write_outputs(lines, summary, out_dir):
    os.makedirs(out_dir, exist_ok=True)
    with open(os.path.join(out_dir, "t12_tagged.csv"), "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["account", "name", "section", "total_12mo", "code", "model_line", "tagged_by"])
        for ln in lines:
            w.writerow([ln.account, ln.name, ln.section, round(ln.total, 2), ln.code or "",
                        MODEL_CODES.get(ln.code, (None, ""))[1], ln.tagged_by or "UNTAGGED"])
    with open(os.path.join(out_dir, "summary.json"), "w") as f:
        json.dump(summary, f, indent=2)

    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "T-12"
    ws.append(["Account"] + [f"M{i}" for i in range(1, 13)] + ["12-Month Total", "Code"])
    for ln in lines:
        months = (ln.months + [0.0] * 12)[:12]
        ws.append([f"{ln.account} - {ln.name}".strip(" -")] + months + [ln.total, ln.code or ""])
    wb.save(os.path.join(out_dir, "t12_for_model.xlsx"))


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("t12")
    ap.add_argument("--sheet")
    ap.add_argument("--map", default=os.path.join(os.path.dirname(__file__), "..", "..", "mappings", "onesite_gl_map.csv"))
    ap.add_argument("--no-map", action="store_true", help="keyword rules only (tests how well rules generalize)")
    ap.add_argument("--out", default="out")
    a = ap.parse_args(argv)

    lines, noi, _ = parse_t12(_read_grid(a.t12, a.sheet))
    tag_lines(lines, {} if a.no_map else load_map(a.map))
    s = summarize(lines, noi)
    write_outputs(lines, s, a.out)

    print(f"Lines: {s['lines']}  tagged: {s['tagged']}  untagged (non-zero): {len(s['untagged_nonzero'])}")
    for u in s["untagged_nonzero"]:
        print(f"  UNTAGGED  {u['account']}  {u['name']}  {u['total']:,.2f}")
    print(f"T-12 NOI: {s['t12_noi']:,.2f}   model NOI: {s['model_noi_before_reserves']:,.2f}   -> {s['noi_check']}")
    print(f"Outputs written to {a.out}/")
    return 0 if s["noi_check"] == "PASS" and not s["untagged_nonzero"] else 1


if __name__ == "__main__":
    sys.exit(main())
