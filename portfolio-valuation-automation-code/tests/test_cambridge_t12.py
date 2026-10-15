"""Cambridge Court is the answer key: the tagger must reproduce Abraham's T12 column."""
import csv
import os
import sys

import openpyxl
import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "src"))
from valuation.t12 import _read_grid, parse_t12, tag_lines, load_map, summarize  # noqa: E402

DEAL = os.path.join(ROOT, "reference-deal", "cambridge-court")
T12 = os.path.join(DEAL, "Cambridge 3-20-2026 Request.xls")
MODEL = os.path.join(DEAL, "CA - Cambridge Court_v4_AR.xlsx")
GL_MAP = os.path.join(ROOT, "mappings", "onesite_gl_map.csv")
T12_NOI = 986818.73


def run(use_map=True):
    lines, noi, _ = parse_t12(_read_grid(T12, "T-12"))
    tag_lines(lines, load_map(GL_MAP) if use_map else {})
    return lines, summarize(lines, noi)


def test_parses_every_line_and_noi():
    lines, s = run()
    assert s["t12_noi"] == pytest.approx(T12_NOI)
    assert len(lines) == 78


@pytest.mark.parametrize("use_map", [True, False], ids=["gl_map", "keywords_only"])
def test_all_lines_tagged_and_noi_ties(use_map):
    _, s = run(use_map)
    assert s["untagged_nonzero"] == []
    assert s["noi_check"] == "PASS", s["noi_difference"]


def test_matches_abrahams_t12_column():
    """Every Scenario 0 T12 value matches the Cambridge model, except:
    RES  - comes from the Dashboard, not the T-12
    OFF  - the model left 'Paymode Rebates' (-$352.04) untagged; we tag it to OFF
    """
    _, s = run()
    ws = openpyxl.load_workbook(MODEL, data_only=True)["Scenario 0"]
    expected_diffs = {"RES": None, "OFF": -352.04}
    for r in range(7, 80):
        code = ws.cell(r, 1).value
        if not code:
            continue
        ca = ws.cell(r, 5).value or 0
        ours = s["scenario0_t12_column"].get(code, 0)
        if code in expected_diffs:
            if expected_diffs[code] is not None:
                assert ours - ca == pytest.approx(expected_diffs[code], abs=0.01)
            continue
        assert ours == pytest.approx(ca, abs=0.5), code


def test_keyword_rules_agree_with_gl_map():
    """Keyword rules alone should agree with Abraham's tags on Cambridge.
    One known, intentional difference: 'Electricity Vacant' -> VAC (model has a
    Vacant Electric row); Cambridge put it in ELEC."""
    key = {r["account"]: r["code"] for r in csv.DictReader(open(GL_MAP))}
    lines, _ = run(use_map=False)
    diffs = {ln.account for ln in lines if ln.account in key and key[ln.account] != ln.code}
    assert diffs == {"6430-0000"}
